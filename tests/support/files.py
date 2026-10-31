"""Helpers that write small input files for reading tests."""
import datetime as dt

import openpyxl


def write_bytes(tmp_path, name, data):
    p = tmp_path / name
    p.write_bytes(data)
    return p


def write_text(tmp_path, name, text, encoding="utf-8"):
    return write_bytes(tmp_path, name, text.encode(encoding))


def write_xlsx(tmp_path, name, sheets):
    """sheets: {sheet name: [[value or (value, number_format), ...], ...]}. A sheet name starting with '~' is hidden."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        hidden = title.startswith("~")
        ws = wb.create_sheet(title.lstrip("~"))
        if hidden:
            ws.sheet_state = "hidden"
        for r, row in enumerate(rows, start=1):
            for c, v in enumerate(row, start=1):
                fmt = None
                if isinstance(v, tuple):
                    v, fmt = v
                if v is None:
                    continue
                cell = ws.cell(row=r, column=c, value=v)
                if fmt:
                    cell.number_format = fmt
    p = tmp_path / name
    wb.save(p)
    return p


DATE = dt.date
