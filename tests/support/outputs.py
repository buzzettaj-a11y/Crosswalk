"""Helpers for M09 tests: download into a folder and read the files back."""
import csv
import io
import json
import pathlib

from wholecount.engine.outputs.view import ExportSettings
from wholecount.session.export import download


def write(s, folder, **settings):
    folder = pathlib.Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    return download(s, folder, ExportSettings(**settings), overwrite=True, confirmed_synced=True)


def read(path):
    """[{column: value}] for a CSV, JSON or Excel file."""
    path = pathlib.Path(path)
    if path.suffix == ".csv":
        return list(csv.DictReader(io.StringIO(path.read_bytes().decode("utf-8-sig"))))
    if path.suffix == ".json":
        return json.loads(path.read_text(encoding="utf-8"))
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb["Data"]
    rows = [[c.value or "" for c in r] for r in ws.iter_rows()]
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def header(path):
    path = pathlib.Path(path)
    if path.suffix == ".csv":
        return next(csv.reader(io.StringIO(path.read_bytes().decode("utf-8-sig"))))
    if path.suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        return list(data[0]) if data else []
    import openpyxl
    ws = openpyxl.load_workbook(path)["Data"]
    return [c.value for c in next(ws.iter_rows())]


def all_bytes(folder):
    """Every file's text, Excel files unpacked, for value scans."""
    import zipfile
    out = {}
    for p in pathlib.Path(folder).rglob("*"):
        if p.is_file():
            if p.suffix == ".xlsx":
                with zipfile.ZipFile(p) as z:
                    out[p.name] = "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist())
            else:
                out[p.name] = p.read_bytes().decode("utf-8", "replace")
    return out
