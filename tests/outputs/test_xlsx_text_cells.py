"""OUT-12: Excel cells are text, so =1+1 stays text and leading zeros stay. JSON values are all strings."""
import openpyxl

from wholecount.engine.outputs.writer import encode_json, encode_xlsx


def test_text_cells(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(encode_xlsx(["A", "B"], [{"A": "=1+1", "B": "00123"}, {"A": "@x", "B": "011"}]))
    ws = openpyxl.load_workbook(p)["Data"]
    assert [[c.value for c in r] for r in ws.iter_rows()] == [["A", "B"], ["=1+1", "00123"], ["@x", "011"]]
    assert all(c.data_type == "s" for r in ws.iter_rows() for c in r)
    import json
    assert json.loads(encode_json(["A"], [{"A": 5}])) == [{"A": "5"}]


def test_xlsx_is_byte_stable():
    a = encode_xlsx(["A"], [{"A": "1"}])
    assert a == encode_xlsx(["A"], [{"A": "1"}])
