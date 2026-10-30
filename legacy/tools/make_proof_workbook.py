"""
Proof workbook: one sheet per thing the tool does, with rows that show what is
and is not affected, and an Expected column stating the outcome.

Two uses. A person loads Proof_Workbook.xlsx into the tool and checks that
each row lands where its Expected column says. And the test suite runs every
sheet through the engine and checks the same thing automatically, so the
workbook can never drift from what the tool actually does.

The sheets are defined once, here. Build the workbook with:

    python tools/make_proof_workbook.py

Verify it (what tests/test_proof.py does):

    python tools/make_proof_workbook.py --verify
"""
from __future__ import annotations

import os
import sys
import warnings

warnings.filterwarnings("ignore")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.chdir(ROOT)

import pandas as pd

import app as engine
import identity
import program_names
import table_detect

OUT = os.path.join(ROOT, "proof", "Proof_Workbook.xlsx")

# A base row that passes every check, so each sheet changes one thing at a time.
BASE = {"client_id": "", "first_name": "Ada", "middle_initial": "",
        "last_name": "Sample", "address_1": "209 Joralemon St", "address_2": "",
        "city": "Brooklyn", "state": "NY", "zip": "11201", "program": "Youth"}

STD_COLS = ["client_id", "first_name", "middle_initial", "last_name",
            "address_1", "address_2", "city", "state", "zip", "program"]


def row(expected, want=None, **over):
    r = dict(BASE)
    r.update(over)
    r["Expected"] = expected
    r["_want"] = want or {}
    return r


PEOPLE = [("Ada", "Sample"), ("Ben", "Fixture"), ("Cara", "Example"),
          ("Devi", "Mock"), ("Evan", "Draft"), ("Fay", "Proxy"),
          ("Gus", "Stub"), ("Hana", "Trial"), ("Ivo", "Dummy")]


def ids(rows):
    """
    Give every row its own identifier and its own person, unless the sheet set
    them on purpose. Rows that share a name and address would otherwise be
    marked as a possible duplicate of each other, which is a real behavior but
    not the one the sheet is showing.
    """
    for i, r in enumerate(rows, start=1):
        if r.get("client_id", "") == "":
            r["client_id"] = f"P{i:03d}"
        if r.get("first_name") == BASE["first_name"] and \
                r.get("last_name") == BASE["last_name"]:
            fn, ln = PEOPLE[(i - 1) % len(PEOPLE)]
            r["first_name"], r["last_name"] = fn, ln
        if r.get("full_name") is None and "full_address" in r and \
                r.get("address_1") == BASE["address_1"]:
            pass
    return rows


# --------------------------------------------------------------------------
# The sheets. Expected starts with one of: CLEAN, REVIEW, SET ASIDE,
# POSSIBLE DUPLICATE. The words after it say what to look for.
# --------------------------------------------------------------------------
SHEETS = []


def sheet(name, note, rows, cols=None, kind="pipeline", extra=None):
    SHEETS.append({"name": name, "note": note, "rows": rows,
                   "cols": cols or STD_COLS, "kind": kind, "extra": extra or {}})


sheet("01 Junk and spacing",
      "Placeholders that mean nothing are emptied. Stray spaces and hidden "
      "characters are removed. Real values are left exactly as entered.",
      ids([
          row("CLEAN: first name 'NULL' is emptied, so the row is flagged as missing it",
              first_name="NULL"),
          row("CLEAN: '#N/A' in Address 2 is emptied", address_2="#N/A",
              want={"Address_2": ""}),
          row("CLEAN: double space and trailing space removed from the street",
              address_1="209  Joralemon St ", want={"Address_1": "209 Joralemon St"}),
          row("CLEAN: a non-breaking space in the name becomes an ordinary space",
              first_name="Mary\u00a0Beth", want={"First_Name": "Mary Beth"}),
          row("CLEAN: nothing to clean, left as entered", first_name="Lena",
              want={"First_Name": "Lena"}),
          row("CLEAN: an apostrophe is real punctuation and is kept",
              last_name="O'Brien", want={"Last_Name": "O'Brien"}),
          row("CLEAN: a hyphen in a name is kept", last_name="Smith-Jones",
              want={"Last_Name": "Smith-Jones"}),
      ]))
# the first row above is actually a REVIEW: NULL empties a required field
SHEETS[-1]["rows"][0]["Expected"] = ("REVIEW: 'NULL' is a placeholder, so the first "
                                     "name is treated as blank and the row is missing it")

sheet("02 Name casing",
      "Names are kept exactly as entered unless you add a tidy-up step: on Tidy "
      "up, add Clean columns first_name and last_name with Smart Title (names). "
      "With that step, only a value entirely in upper case or entirely in lower "
      "case is re-cased; mixed case is treated as deliberate and kept.",
      ids([
          row("CLEAN: JOHN becomes John", first_name="JOHN", want={"First_Name": "John"}),
          row("CLEAN: smith becomes Smith", last_name="smith", want={"Last_Name": "Smith"}),
          row("CLEAN: MCMAHON becomes McMahon", last_name="MCMAHON",
              want={"Last_Name": "McMahon"}),
          row("CLEAN: O'BRIEN becomes O'Brien", last_name="O'BRIEN",
              want={"Last_Name": "O'Brien"}),
          row("CLEAN: DiNapoli is mixed case, kept exactly", last_name="DiNapoli",
              want={"Last_Name": "DiNapoli"}),
          row("CLEAN: LaGuardia is mixed case, kept exactly", last_name="LaGuardia",
              want={"Last_Name": "LaGuardia"}),
          row("CLEAN: mELINDA is mixed case, so it is NOT corrected; fix by hand",
              first_name="mELINDA", want={"First_Name": "mELINDA"}),
      ]), kind="casing")

sheet("03 Middle initial",
      "A middle name or initial is reduced to one upper case letter.",
      ids([
          row("CLEAN: James becomes J", middle_initial="James", want={"Middle_Initial": "J"}),
          row("CLEAN: j. becomes J", middle_initial="j.", want={"Middle_Initial": "J"}),
          row("CLEAN: a becomes A", middle_initial="a", want={"Middle_Initial": "A"}),
          row("CLEAN: already J, unchanged", middle_initial="J", want={"Middle_Initial": "J"}),
          row("CLEAN: blank stays blank", middle_initial="", want={"Middle_Initial": ""}),
      ]))

sheet("04 ZIP forms",
      "A ZIP is reduced to five digits. Anything that is not five digits sends "
      "the row to review; it is never padded or guessed.",
      ids([
          row("CLEAN: ZIP+4 11201-1234 becomes 11201", zip="11201-1234", want={"Zip": "11201"}),
          row("CLEAN: spaces around 11201 removed", zip=" 11201 ", want={"Zip": "11201"}),
          row("CLEAN: 11201 as entered", zip="11201", want={"Zip": "11201"}),
          row("REVIEW: four digits 7030 is not a ZIP; a leading zero was probably lost",
              zip="7030"),
          row("REVIEW: letters ABCDE are not a ZIP", zip="ABCDE"),
          row("REVIEW: blank ZIP", zip=""),
          row("CLEAN: 112011234 is ZIP+4 without the hyphen, trimmed to 11201",
              zip="112011234", want={"Zip": "11201"}),
      ]))

sheet("05 City from ZIP",
      "A blank city is filled from a New York City ZIP. A city that is present "
      "is never overwritten by this rule.",
      ids([
          row("CLEAN: blank city, ZIP 11201 gives Brooklyn", city="", zip="11201",
              want={"City": "Brooklyn"}),
          row("CLEAN: blank city, ZIP 10007 gives New York", city="", zip="10007",
              want={"City": "New York"}),
          row("CLEAN: blank city, ZIP 10306 gives Staten Island", city="", zip="10306",
              want={"City": "Staten Island"}),
          row("CLEAN: blank city, ZIP 11377 gives Queens", city="", zip="11377",
              want={"City": "Queens"}),
          row("CLEAN: city present, kept", city="Brooklyn", want={"City": "Brooklyn"}),
          row("REVIEW: blank city and ZIP 07030 is not New York City, so nothing fills it",
              city="", zip="07030"),
      ]))

sheet("06 City against ZIP",
      "A recognized borough that disagrees with its ZIP is a contradiction and "
      "goes to review showing both. An unrecognized city with a valid NYC ZIP "
      "is replaced by the ZIP's borough and recorded, because the ZIP is the "
      "machine-generated one of the two.",
      ids([
          row("REVIEW: Brooklyn with a Manhattan ZIP 10007", city="Brooklyn", zip="10007"),
          row("REVIEW: Queens with a Brooklyn ZIP 11201", city="Queens", zip="11201"),
          row("REVIEW: Bronx with a Staten Island ZIP 10306", city="Bronx", zip="10306"),
          row("CLEAN: Brooklyn with 11201 agree", city="Brooklyn", zip="11201"),
          row("CLEAN: New York with 10007 agree", city="New York", zip="10007"),
          row("CLEAN: Manhattan is another name for New York, agrees with 10007",
              city="Manhattan", zip="10007", want={"City": "New York"}),
          row("CLEAN: Springfield is not recognized, replaced by Brooklyn from ZIP 11201 and recorded",
              city="Springfield", zip="11201", want={"City": "Brooklyn"}),
      ]))

sheet("07 Borough names",
      "Abbreviations and neighborhoods are turned into the borough they belong to.",
      ids([
          row("CLEAN: Bklyn becomes Brooklyn", city="Bklyn", want={"City": "Brooklyn"}),
          row("CLEAN: BK becomes Brooklyn", city="BK", want={"City": "Brooklyn"}),
          row("CLEAN: Woodside is in Queens", city="Woodside", zip="11377",
              want={"City": "Queens"}),
          row("CLEAN: Bay Ridge is in Brooklyn", city="Bay Ridge", zip="11209",
              want={"City": "Brooklyn"}),
          row("CLEAN: Harlem is in New York", city="Harlem", zip="10027",
              want={"City": "New York"}),
          row("CLEAN: Brooklyn already, unchanged", city="Brooklyn", want={"City": "Brooklyn"}),
      ]))

sheet("08 State",
      "State is written as two letters. Blank becomes NY. Anything other than "
      "New York goes to review, because this is a New York City dataset.",
      ids([
          row("CLEAN: blank becomes NY", state="", want={"State": "NY"}),
          row("CLEAN: 'New York' becomes NY", state="New York", want={"State": "NY"}),
          row("CLEAN: 'ny' becomes NY", state="ny", want={"State": "NY"}),
          row("CLEAN: NY as entered", state="NY", want={"State": "NY"}),
          row("REVIEW: NJ is not New York", state="NJ"),
          row("REVIEW: CT is not New York", state="CT"),
      ]))

sheet("09 Jammed address",
      "One column holding street, unit, city, state and ZIP together is split "
      "into its parts. Streets that merely contain a unit word are left alone.",
      ids([
          row("CLEAN: split into 123 Main St, Apt 3L, Brooklyn, 11201",
              full_address="123 Main St Apt 3L Brooklyn NY 11201",
              want={"Address_1": "123 Main St", "Address_2": "Apt 3L",
                    "City": "Brooklyn", "Zip": "11201"}),
          row("CLEAN: commas: 45 Oak St, Queens, 11375",
              full_address="45 Oak St, Queens, NY 11375",
              want={"Address_1": "45 Oak St", "City": "Queens", "Zip": "11375"}),
          row("CLEAN: dashes as separators: unit Apt 2 comes out",
              full_address="60 Centre St - Apt 2 - New York NY 10007",
              want={"Address_1": "60 Centre St", "Address_2": "Apt 2", "Zip": "10007"}),
          row("CLEAN: trailing USA does not hide the ZIP",
              full_address="100 Grand St Apt 4 Brooklyn NY 11211 USA",
              want={"Address_1": "100 Grand St", "Zip": "11211"}),
          row("CLEAN: 88 Front St is a street, Front is not a unit here",
              full_address="88 Front St Brooklyn NY 11201",
              want={"Address_1": "88 Front St", "Address_2": ""}),
          row("CLEAN: Grand Concourse is a street, left whole",
              full_address="2488 Grand Concourse Bronx NY 10458",
              want={"Address_1": "2488 Grand Concourse", "Address_2": ""}),
          row("CLEAN: 40-22 24th St keeps its hyphenated number",
              full_address="40-22 24th St Apt 5 Queens NY 11101",
              want={"Address_1": "40-22 24th St", "Address_2": "Apt 5"}),
          row("CLEAN: unit as its own comma part, city missing: Apt. 4 is the unit, city filled from ZIP",
              full_address="111 Test St, Apt. 4, NY, 11201",
              want={"Address_1": "111 Test St", "Address_2": "Apt. 4", "City": "Brooklyn"}),
          row("CLEAN: unit written after the city is still the unit",
              full_address="111 Test St, Brooklyn, Apt. 4, NY, 11201",
              want={"Address_1": "111 Test St", "Address_2": "Apt. 4", "City": "Brooklyn"}),
          row("CLEAN: 4th Fl is a unit, not a city",
              full_address="111 Test St, 4th Fl, Brooklyn, NY, 11201",
              want={"Address_2": "4th Fl", "City": "Brooklyn"}),
          row("CLEAN: the state spelled out is the state, not the city",
              full_address="111 Test St, Apt. 4, Brooklyn, New York, 11201",
              want={"Address_2": "Apt. 4", "City": "Brooklyn", "State": "NY"}),
          row("CLEAN: a bare unit like 14A between the street and the city is the unit, never the city",
              full_address="117-19 45TH AVE, 14A, QUEENS, NY 11355",
              want={"Address_1": "117-19 45TH AVE", "Address_2": "14A", "City": "Queens", "Zip": "11355"}),
          row("CLEAN: 14A jammed onto the street with no comma is still the unit",
              full_address="117-19 45TH AVE 14A QUEENS NY 11355",
              want={"Address_1": "117-19 45TH AVE", "Address_2": "14A", "City": "Queens"}),
          row("CLEAN: Avenue U is a street name, not a unit",
              full_address="2300 Avenue U, Brooklyn, NY 11229",
              want={"Address_1": "2300 Avenue U", "Address_2": ""}),
      ]),
      cols=["client_id", "first_name", "last_name", "full_address", "program"])

sheet("10 Jammed name",
      "One column holding the whole name is split into first, middle and last. "
      "Both orders are handled, a suffix stays with the surname, and casing "
      "such as DiNapoli is kept.",
      ids([
          row("CLEAN: 'Smith, John A' gives John A Smith", full_name="Smith, John A",
              want={"First_Name": "John", "Middle_Initial": "A", "Last_Name": "Smith"}),
          row("CLEAN: 'John A Smith' gives the same", full_name="John A Smith",
              address_1="60 Centre St", city="New York", zip="10007",
              want={"First_Name": "John", "Middle_Initial": "A", "Last_Name": "Smith"}),
          row("CLEAN: 'Smith Jr, John' keeps Jr with the surname", full_name="Smith Jr, John",
              want={"First_Name": "John", "Last_Name": "Smith Jr"}),
          row("CLEAN: 'DiNapoli, Anthony' keeps DiNapoli's casing",
              full_name="DiNapoli, Anthony", want={"Last_Name": "DiNapoli"}),
          row("CLEAN: 'MARIA GARCIA' splits but keeps its case; add Smart Title to re-case",
              full_name="MARIA GARCIA", want={"First_Name": "MARIA", "Last_Name": "GARCIA"}),
          row("CLEAN: 'Ana Lopez' has no middle initial", full_name="Ana Lopez",
              want={"First_Name": "Ana", "Middle_Initial": "", "Last_Name": "Lopez"}),
      ]),
      cols=["client_id", "full_name", "address_1", "city", "state", "zip", "program"])

sheet("11 Unit in both lines",
      "When the street also holds a unit, it depends what Address 2 has. The "
      "same unit: the copy in the street is removed. A different unit: review, "
      "showing both. Something that is not a unit: review, nothing overwritten.",
      ids([
          row("CLEAN: Apt 2A in both, the street copy is removed",
              address_1="123 Main St Apt 2A", address_2="Apt 2A",
              want={"Address_1": "123 Main St", "Address_2": "Apt 2A"}),
          row("CLEAN: #2A and Apt 2A are the same unit",
              address_1="123 Main St Apt 2A", address_2="#2A",
              want={"Address_1": "123 Main St"}),
          row("CLEAN: unit only in the street, moved to Address 2",
              address_1="123 Main St Apt 2A", address_2="",
              want={"Address_1": "123 Main St", "Address_2": "Apt 2A"}),
          row("CLEAN: unit only in Address 2, nothing to do",
              address_1="123 Main St", address_2="Apt 2A",
              want={"Address_1": "123 Main St", "Address_2": "Apt 2A"}),
          row("CLEAN: a space inside the designator, Apt. 10 R, is still one unit",
              address_1="67 Franklin Avenue Apt. 10 R", address_2="",
              want={"Address_1": "67 Franklin Avenue", "Address_2": "Apt. 10 R"}),
          row("CLEAN: APT18L jammed onto the word is a unit",
              address_1="25 Tom Street APT18L", address_2="",
              want={"Address_1": "25 Tom Street", "Address_2": "APT18L"}),
          row("CLEAN: Apt#: 11T, with a hash and colon, is a unit",
              address_1="67 Jill Ave Apt#: 11T", address_2="",
              want={"Address_1": "67 Jill Ave", "Address_2": "Apt#: 11T"}),
          row("CLEAN: Apartment Way is a street name, not a unit",
              address_1="10 Apartment Way", address_2="",
              want={"Address_1": "10 Apartment Way", "Address_2": ""}),
          row("REVIEW: Apt 2A in the street but Apt 3B in Address 2, a contradiction",
              address_1="123 Main St Apt 2A", address_2="Apt 3B"),
          row("REVIEW: Address 2 holds 'c/o Jane Doe', not a unit, so a person decides",
              address_1="123 Main St Apt 2A", address_2="c/o Jane Doe"),
      ]))

sheet("12 Missing required",
      "A required field left blank sends the row to review. Optional fields "
      "may be blank.",
      ids([
          row("REVIEW: no first name", first_name=""),
          row("REVIEW: no last name", last_name=""),
          row("REVIEW: no street address", address_1=""),
          row("REVIEW: no ZIP", zip=""),
          row("CLEAN: middle initial is optional", middle_initial=""),
          row("CLEAN: Address 2 is optional", address_2=""),
          row("CLEAN: program is optional", program=""),
      ]))

sheet("13 Digits in a name",
      "A name containing digits usually means the columns are matched one "
      "across, so the row is sent to review rather than cleaned.",
      ids([
          row("REVIEW: J0hn has a zero in it", first_name="J0hn"),
          row("REVIEW: a house number in the surname column", last_name="209 Joralemon"),
          row("REVIEW: 3rd in a first name", first_name="3rd"),
          row("CLEAN: Mary-Kate has no digits", first_name="Mary-Kate"),
          row("CLEAN: St. John has no digits", last_name="St. John"),
          row("CLEAN: Smith III, a suffix, is fine", last_name="Smith III"),
      ]))

sheet("14 Same identifier",
      "Rows sharing one identifier are either identical, in which case one is "
      "kept and the rest set aside, or different, in which case all of them go "
      "to review and nothing is merged.",
      [
          row("CLEAN: first of three identical rows with ID D100, kept", client_id="D100"),
          row("SET ASIDE: identical copy of D100", client_id="D100"),
          row("SET ASIDE: identical copy of D100", client_id="D100"),
          row("REVIEW: ID D200 also used by a row at a different address", client_id="D200",
              first_name="Ben", last_name="Fixture"),
          row("REVIEW: same ID D200, different address, never merged", client_id="D200",
              first_name="Ben", last_name="Fixture",
              address_1="60 Centre St", city="New York", zip="10007"),
          row("CLEAN: ID D300 appears once", client_id="D300", first_name="Cara",
              last_name="Example", address_1="87 Bay Terrace",
              city="Staten Island", zip="10306"),
      ])

sheet("15 Same person twice in one file",
      "Rows in one file that match on First_Name, Last_Name and Zip are "
      "proposed on Tidy up, shown side by side, and decided by a person. "
      "Nothing is merged and nothing is removed until it is confirmed. This "
      "sheet confirms the first group as copies and leaves the rest undecided.",
      [
          row("CLEAN: first of a confirmed pair, kept", client_id="Q1",
              first_name="Maria", last_name="Garcia", address_1="330 Frost St",
              address_2="Apt 7D", zip="11221"),
          row("SET ASIDE: confirmed as a copy of the row above", client_id="Q2",
              first_name="Maria", last_name="Garcia", address_1="330 Frost St",
              address_2="Apt 7D", zip="11221"),
          row("CLEAN: same name at a different ZIP is never proposed", client_id="Q4",
              first_name="Maria", last_name="Garcia", address_1="1 Bowling Green",
              city="New York", zip="10004"),
          row("CLEAN: proposed with the next row but not decided, so kept and flagged",
              client_id="Q5", first_name="Devi", last_name="Mock",
              address_1="87 Bay Terrace", city="Staten Island", zip="10306"),
          row("CLEAN: proposed with the row above but not decided, so kept and flagged",
              client_id="Q6", first_name="Devi", last_name="Mock",
              address_1="87 Bay Terrace", city="Staten Island", zip="10306"),
      ], kind="tidyup_first")

sheet("16 Program names",
      "The same program written differently is grouped on the matching stage "
      "and counted under its most used spelling. Rename the group there if you "
      "want another name. Different programs stay separate.",
      ids([
          row("CLEAN: counted as Workforce, the most used spelling", program="Workforce",
              want={"__program__": "Workforce"}),
          row("CLEAN: counted as Workforce", program="Workforce",
              want={"__program__": "Workforce"}),
          row("CLEAN: Workforce Development is the same program", program="Workforce Development",
              want={"__program__": "Workforce"}),
          row("CLEAN: WFD is its initials", program="WFD", want={"__program__": "Workforce"}),
          row("CLEAN: Resident Services stands alone", program="Resident Services",
              want={"__program__": "Resident Services"}),
          row("CLEAN: Resident Svcs is an abbreviation of it", program="Resident Svcs",
              want={"__program__": "Resident Services"}),
          row("CLEAN: Homelessness Prevention is a different program", program="Homelessness Prevention",
              want={"__program__": "Homelessness Prevention"}),
      ]), kind="programs")

sheet("17 Field types",
      "Add three columns on Set up your columns first: Enrolled as a date, "
      "Hours as a decimal, Active as yes/no. A value that fits is written in "
      "one standard form. One that does not fit sends the row to review.",
      ids([
          row("CLEAN: 3/15/2024 stored as 2024-03-15", enrolled="3/15/2024", hours="12.5",
              active="Y", want={"c:Enrolled": "2024-03-15", "c:Active": "Yes"}),
          row("CLEAN: 2024-03-15 already standard", enrolled="2024-03-15", hours="8",
              active="no", want={"c:Enrolled": "2024-03-15", "c:Active": "No"}),
          row("CLEAN: 15.03.2024 read day first", enrolled="15.03.2024", hours="1,200.75",
              active="true", want={"c:Enrolled": "2024-03-15", "c:Hours": "1200.75"}),
          row("REVIEW: 'not a date' is not a date", enrolled="not a date", hours="8",
              active="Y"),
          row("REVIEW: 'x' is not a number", enrolled="2024-03-15", hours="x", active="Y"),
          row("REVIEW: 'maybe' is not yes or no", enrolled="2024-03-15", hours="8",
              active="maybe"),
      ]),
      cols=STD_COLS + ["enrolled", "hours", "active"], kind="types")

sheet("18 Same person across systems",
      "Seen on the Results stage. Add a column named DOB as a date on Set up "
      "your columns. Pairs that may be one person are proposed for a decision; "
      "different people are not.",
      [
          row("PROPOSED with row 2: surname changed, same birth date and address",
              client_id="S1", first_name="Maria", last_name="Delgado", dob="1988-04-11",
              address_1="209 Joralemon St", zip="11201"),
          row("PROPOSED with row 1", client_id="S2", first_name="Maria", last_name="Okafor",
              dob="1988-04-11", address_1="209 Joralemon St", zip="11201",
              program="Workforce"),
          row("PROPOSED with row 4: Bob is Robert", client_id="S3", first_name="Robert",
              last_name="Newman", dob="1970-09-14", address_1="88 Meserole St",
              city="Brooklyn", zip="11206"),
          row("PROPOSED with row 3", client_id="S4", first_name="Bob", last_name="Newman",
              dob="1970-09-14", address_1="88 Meserole St", city="Brooklyn",
              zip="11206", program="Workforce"),
          row("PROPOSED with row 6: day and month swapped in the birth date",
              client_id="S5", first_name="Samuel", last_name="Cho", dob="1999-01-02",
              address_1="60 Centre St", city="New York", zip="10007"),
          row("PROPOSED with row 5", client_id="S6", first_name="Samuel", last_name="Cho",
              dob="1999-02-01", address_1="60 Centre St", city="New York",
              zip="10007", program="Workforce"),
          row("NOT PROPOSED: a different person at the same address",
              client_id="S7", first_name="Tomas", last_name="Reyes", dob="1975-01-02",
              address_1="209 Joralemon St", zip="11201"),
      ],
      cols=STD_COLS + ["dob"], kind="identity")

sheet("21 Duplicates within a file",
      "Tidy up proposes rows that match on First_Name, Last_Name and Zip within "
      "one file and shows them side by side. Confirm each group as copies to "
      "have the first kept and the rest set aside; nothing is removed until you "
      "do. Matching ignores case, spacing and punctuation, and a ZIP+4 is its "
      "first five digits. A row blank in a chosen field is never proposed.",
      [
          row("CLEAN: first of a pair, kept and marked as group F1", client_id="U1",
              first_name="Ada", last_name="Sample", zip="11201"),
          row("SET ASIDE: ADA sample at 11201-1234 is the same by the rule, a copy of the row above",
              client_id="U2", first_name="ADA", last_name="sample", zip="11201-1234"),
          row("CLEAN: same name at a different ZIP is not matched", client_id="U3",
              first_name="Ada", last_name="Sample", address_1="60 Centre St",
              city="New York", zip="10007"),
          row("CLEAN: first of a second pair, group F2", client_id="U4",
              first_name="Ben", last_name="Fixture", address_1="88 Meserole St",
              zip="11206"),
          row("SET ASIDE: ben FIXTURE at 11206, a copy of the row above", client_id="U5",
              first_name="ben", last_name="FIXTURE", address_1="88 Meserole Street",
              zip="11206"),
          row("CLEAN: nobody else matches Cara Example", client_id="U6",
              first_name="Cara", last_name="Example", address_1="87 Bay Terrace",
              city="Staten Island", zip="10306"),
      ], kind="rule")

# Table detection sheets are laid out as grids, not rows.
sheet("19 Two tables, one sheet",
      "A summary table sits beside the client list with a blank column between. "
      "The tool should list two tables and include only the client list.",
      [], kind="grid",
      extra={"grid": [
          ["Program", "Total", "", "client_id", "first_name", "last_name",
           "address_1", "city", "state", "zip"],
          ["Youth", "124", "", "G1", "Ada", "Sample", "209 Joralemon St",
           "Brooklyn", "NY", "11201"],
          ["Workforce", "87", "", "G2", "Ben", "Fixture", "60 Centre St",
           "New York", "NY", "10007"],
          ["Housing", "206", "", "G3", "Cara", "Example", "87 Bay Terrace",
           "Staten Island", "NY", "10306"],
      ], "expect_tables": 2})

sheet("20 Title rows above header",
      "Two title rows sit above the real header. Choose row 3 as the header on "
      "the matching stage and the columns match; the title rows are dropped.",
      [], kind="grid",
      extra={"grid": [
          ["Client roster, spring 2026", "", "", "", "", "", ""],
          ["Prepared by the data team", "", "", "", "", "", ""],
          ["client_id", "first_name", "last_name", "address_1", "city", "state", "zip"],
          ["H1", "Ada", "Sample", "209 Joralemon St", "Brooklyn", "NY", "11201"],
          ["H2", "Ben", "Fixture", "60 Centre St", "New York", "NY", "10007"],
          ["H3", "Cara", "Example", "87 Bay Terrace", "Staten Island", "NY", "10306"],
      ], "header_row": 3})


# --------------------------------------------------------------------------
# Building the workbook
# --------------------------------------------------------------------------
def build():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    intro = [
        ["Crosswalk proof workbook"],
        [""],
        ["Each sheet shows one thing the tool does. Load this workbook, match the "
         "columns as usual, process it, and check every row lands where its "
         "Expected column says."],
        [""],
        ["Expected begins with one of:"],
        ["  CLEAN", "the row passes; anything after the colon is a change to look for"],
        ["  REVIEW", "the row goes to Fix problems with that reason"],
        ["  SET ASIDE", "an identical copy, kept out of the clean list"],
        ["  POSSIBLE DUPLICATE", "marked and shown, never merged"],
        ["  PROPOSED / NOT PROPOSED", "on the Results stage, whether a pair is offered as one person"],
        [""],
        ["The Expected column is not part of the data. Leave it unmatched."],
        [""],
        ["Sheets 17 and 18 need extra columns on Set up your columns; each sheet says which."],
        ["Sheets 19 and 20 test how tables are found and are checked on the first two stages."],
        [""],
        ["The same expectations are checked automatically by tests/test_proof.py, so this "
         "workbook always matches what the tool does."],
    ]
    for line in intro:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 80

    head_fill = PatternFill("solid", fgColor="1F3B63")
    head_font = Font(bold=True, color="FFFFFF")
    fills = {"CLEAN": "E2EFE4", "REVIEW": "FBE9D0", "SET ASIDE": "F8DEDA",
             "POSSIBLE DUPLICATE": "E8F0F7", "PROPOSED": "E8F0F7",
             "NOT PROPOSED": "E2EFE4"}

    for s in SHEETS:
        w = wb.create_sheet(s["name"][:31])
        w.append([s["note"]])
        w.merge_cells(start_row=1, start_column=1, end_row=1,
                      end_column=max(4, len(s["cols"]) + 1))
        w["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        w.row_dimensions[1].height = 48
        w.append([])
        if s["kind"] == "grid":
            for r in s["extra"]["grid"]:
                w.append(r)
            continue
        w.append(s["cols"] + ["Expected"])
        for c in w[3]:
            c.fill, c.font = head_fill, head_font
        for r in s["rows"]:
            w.append([r.get(c, "") for c in s["cols"]] + [r["Expected"]])
            cell = w.cell(row=w.max_row, column=len(s["cols"]) + 1)
            key = next((k for k in fills if r["Expected"].startswith(k)), None)
            if key:
                cell.fill = PatternFill("solid", fgColor=fills[key])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for i, col in enumerate(s["cols"], start=1):
            w.column_dimensions[w.cell(row=3, column=i).column_letter].width = \
                max(12, min(34, max(len(str(r.get(col, ""))) for r in s["rows"]) + 2))
        w.column_dimensions[w.cell(row=3, column=len(s["cols"]) + 1).column_letter].width = 70
        w.freeze_panes = "A4"

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    return OUT


# --------------------------------------------------------------------------
# Verifying every expectation against the engine
# --------------------------------------------------------------------------
def _schema(extra_fields=()):
    sch = [dict(f) for f in engine.DEFAULT_SCHEMA]
    for f in sch:
        if f["role"] == "Client_ID":
            f["required"] = False
    for label, dtype in extra_fields:
        sch.append({"role": "custom", "label": label, "dtype": dtype,
                    "length": 0, "required": False})
    return sch


def _frame(s):
    return pd.DataFrame([{c: r.get(c, "") for c in s["cols"]} for r in s["rows"]],
                        dtype=str)


def _outcome(rec):
    errs = " ".join(rec["__errors__"]).lower()
    if "confirmed duplicate: identical" in errs or "by your rule" in errs:
        return "SET ASIDE"
    if rec["__errors__"]:
        return "REVIEW"
    if rec.get("Possible_Duplicate"):
        return "POSSIBLE DUPLICATE"
    return "CLEAN"


def _value(rec, key):
    if key.startswith("c:"):
        return rec.get("__custom__", {}).get(key[2:], "")
    return rec.get(key, "")


def verify():
    results = []

    def check(label, ok):
        results.append((label, bool(ok)))

    for s in SHEETS:
        name = s["name"]
        if s["kind"] == "grid":
            g = s["extra"]["grid"]
            tables = table_detect.detect_in_grid(g, source=name)
            if "expect_tables" in s["extra"]:
                check(f"{name}: {s['extra']['expect_tables']} tables are found",
                      len(tables) == s["extra"]["expect_tables"])
                recs = [t for t in tables if t.kind == "records"]
                check(f"{name}: the client list is the one classed as records",
                      len(recs) == 1 and "first_name" in recs[0].headers)
            if "header_row" in s["extra"]:
                h, rows = table_detect.frame_from_block(g, s["extra"]["header_row"])
                check(f"{name}: choosing row {s['extra']['header_row']} gives the real headers",
                      "client_id" in h and len(rows) == 3)
            continue

        extra = {"types": [("Enrolled", "date"), ("Hours", "decimal"),
                           ("Active", "bool")],
                 "identity": [("DOB", "date")]}.get(s["kind"], [])
        schema = _schema(extra)
        df = _frame(s)
        ops = engine.auto_prep_ops(df)
        if s["kind"] == "casing":
            ops.append({"type": "clean", "columns": ["first_name", "last_name"],
                        "case": "Smart Title (names)", "sanitize": False,
                        "strip_punct": False})
        prep = engine.apply_prep_ops(df, ops)
        mapping = engine.guess_mapping(list(prep.columns))
        if s["kind"] in ("rule", "tidyup_first"):
            import dedupe_rules
            flds = ["First_Name", "Last_Name", "Zip"]
            pr = dedupe_rules.proposals(prep, flds, mapping)
            groups_sorted = sorted(pr["groups"], key=lambda g: g["rows"][0]["row"])
            if s["kind"] == "rule":
                dec = {g["gid"]: "copies" for g in groups_sorted}
            else:
                dec = {groups_sorted[0]["gid"]: "copies"} if groups_sorted else {}
            dedupe_rules.apply_decisions(prep, flds, mapping, dec)
        for label, _t in extra:
            mapping[label] = next((c for c in prep.columns
                                   if c.lower() == label.lower()), "")
        aliases = None
        if s["kind"] == "programs":
            vals = df["program"].value_counts().to_dict()
            aliases = program_names.build_alias_map(
                program_names.suggest_groups(vals))
        res = engine.run_pipeline({name: prep}, {name: mapping}, False, None,
                                  1337, schema=schema, program_aliases=aliases)
        by_row = {int(r["__source_row__"]): r for r in res["transformed"]}

        for i, r in enumerate(s["rows"]):
            rec = by_row.get(i)
            want = r["Expected"].split(":")[0].split(" with ")[0].strip()
            if s["kind"] == "identity":
                continue
            got = _outcome(rec)
            check(f"{name} row {i + 1}: {r['Expected'][:70]}", got == want)
            for key, val in r["_want"].items():
                check(f"{name} row {i + 1}: {key} is {val!r}",
                      _value(rec, key) == val)

        if s["kind"] == "identity":
            recs = [identity.PersonRecord(
                f"r{i}", first=rec["First_Name"], last=rec["Last_Name"],
                dob=rec.get("__custom__", {}).get("DOB", ""),
                street=rec["Address_1"], zip=rec["Zip"],
                program=rec.get("__program__", ""),
                client_id=rec["Client_ID"], system="proof")
                for i, rec in sorted(by_row.items())]
            out = identity.resolve(recs)
            pairs = {frozenset(p["pair"]) for p in out["proposals"]}
            joined = {frozenset(l["pair"]) for l in out["auto_links"]}
            for i, r in enumerate(s["rows"]):
                exp = r["Expected"]
                if exp.startswith("PROPOSED with row"):
                    j = int(exp.split("row")[1].split(":")[0]) - 1
                    pr = frozenset({f"r{i}", f"r{j}"})
                    check(f"{name} row {i + 1}: proposed or joined with row {j + 1}",
                          pr in pairs or pr in joined)
                elif exp.startswith("NOT PROPOSED"):
                    check(f"{name} row {i + 1}: not paired with anyone",
                          not any(f"r{i}" in p for p in pairs | joined))

    failed = [l for l, ok in results if not ok]
    for l in failed:
        print(f"  [FAIL] {l}")
    print(f"PROOF WORKBOOK: {'PASS' if not failed else 'FAIL'} "
          f"({len(results) - len(failed)}/{len(results)})")
    return not failed


if __name__ == "__main__":
    if "--verify" in sys.argv:
        sys.exit(0 if verify() else 1)
    path = build()
    print(f"wrote {path}")
    sys.exit(0 if verify() else 1)
