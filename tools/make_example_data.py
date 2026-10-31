"""Writes the five example files and EXPECTED.json into wholecount/sample_data/ (P-GEO-4a). A dev tool that never ships.

    python tools/make_example_data.py [--seed 20261030] [--out wholecount/sample_data] [--check]

Every person is made up. Names come from fixed lists below and are paired by a seeded random generator; streets are real NYC
street names with plausible house numbers. No value comes from a real person.

It also writes sample_reference.csv, the practice reference for "Try example data" (P-GEO-4, P-GEO-4a). Its first line is
#wholecount-sample, it holds every example address the plan places at a made-up point inside the city, and it leaves out the
few the plan marks as not found. Its SHA-256 must be in SAMPLE_REFERENCE_HASHES in wholecount/engine/geo/reference.py, and
--check fails while it isn't.

EXPECTED.json is written from the plan in this script, never from a run. A milestone that gets a different number finds out why:
if the code is wrong it fixes the code, and if the plan is wrong it changes the plan here, regenerates, and says so in its report.
Nobody edits EXPECTED.json by hand.

Uses the standard library, openpyxl and lxml only.
"""
import argparse
import csv
import hashlib
import datetime as dt
import io
import json
import pathlib
import random
import sys
import tempfile
from dataclasses import dataclass, field

import re

import openpyxl
from lxml import etree

SEED = 20261030
ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "wholecount" / "sample_data"

CSV, XLSX, TSV, JSON_, XML = "example_clinic.csv", "example_housing.xlsx", "example_partner.tsv", "example_portal.json", "example_state.xml"
FILES = (CSV, XLSX, TSV, JSON_, XML)
REFERENCE = "sample_reference.csv"
REFERENCE_PY = ROOT / "wholecount" / "engine" / "geo" / "reference.py"
NOT_FOUND = 3
# Made-up points: a box around a spot inside each borough. Never a real building's location.
BOROUGH_BOX = {"Bronx": (40.845, -73.885), "Brooklyn": (40.650, -73.950), "New York": (40.780, -73.965),
               "Queens": (40.720, -73.830), "Staten Island": (40.600, -74.120)}
# Planned reason IDs whose registry flag is location (M04), so the row is "Counted, not mapped: address problem".
LOCATION_REASONS = {"address_incomplete", "address_part_missing", "zip_not_valid"}
LOCATION_KINDS = ("no_fixed_address", "outside_nyc", "address_problem", "not_found", "no_geocoder")

GIVEN = [
    "Aaliyah", "Adrian", "Aisha", "Alejandro", "Amara", "Andre", "Angela", "Arjun", "Beatriz", "Brandon", "Carmen", "Chen",
    "Chloe", "Claudia", "Damian", "Darnell", "Deborah", "Diego", "Dmitri", "Elena", "Emeka", "Esther", "Fatima", "Felix",
    "Gabriela", "Grace", "Hassan", "Hector", "Imani", "Irene", "Isaac", "Jasmine", "Javier", "Jin", "Joanna", "Jonah", "Kayla",
    "Kenji", "Keisha", "Kofi", "Lamar", "Leah", "Liam", "Lucia", "Malik", "Marcus", "Mariam", "Marisol", "Mei", "Miguel",
    "Mina", "Nadia", "Naomi", "Nikhil", "Noah", "Olga", "Omar", "Priya", "Rafael", "Rania", "Reggie", "Rosa", "Ruth", "Samir",
    "Sasha", "Selena", "Shanice", "Sofia", "Tamara", "Terrence", "Tomas", "Valeria", "Victor", "Wei", "Xavier", "Yara",
    "Yusuf", "Zainab", "Zoe", "Bianca", "Cedric", "Delia", "Ezra", "Fiona", "Gideon", "Hana", "Ivan", "Janelle", "Kwame",
    "Lorena", "Mateo", "Nia", "Oscar", "Paloma", "Quentin", "Renata", "Simone", "Tariq", "Ursula", "Vivian",
]
SURNAME = [
    "Abernathy", "Acosta", "Adeyemi", "Albright", "Alvarado", "Ambrose", "Andersen", "Arroyo", "Bakshi", "Balogun", "Barros",
    "Beaumont", "Benitez", "Blackwood", "Bonilla", "Calloway", "Castellanos", "Chandra", "Chavarria", "Cisneros", "Coleridge",
    "Concepcion", "Corriveau", "Dalrymple", "Danjuma", "Delacroix", "Dempsey", "Duarte", "Easterbrook", "Echeverria", "Efron",
    "Escobedo", "Fairbanks", "Fajardo", "Farouk", "Figueroa", "Fitzwater", "Galloway", "Garibay", "Gonsalves", "Greenleaf",
    "Guerrero", "Halvorsen", "Haverford", "Hinojosa", "Holloway", "Ibarra", "Ilunga", "Iwasaki", "Jaramillo", "Jessup",
    "Kaczmarek", "Kapoor", "Kowalczyk", "Lacroix", "Landry", "Llewellyn", "Lozano", "Madueke", "Malabanan", "Marchetti",
    "Mendoza", "Montague", "Nakashima", "Navarrete", "Nwosu", "Oduya", "Okafor", "Olmstead", "Orellana", "Pacheco", "Pemberton",
    "Quintero", "Rasmussen", "Redmond", "Rinaldi", "Rosales", "Saavedra", "Santangelo", "Sepulveda", "Sokolov", "Stanhope",
    "Tavares", "Thibodeaux", "Toussaint", "Underhill", "Urquhart", "Valdivia", "Vanterpool", "Villanueva", "Wakefield",
    "Whitcombe", "Yamaguchi", "Yeboah", "Zamora", "Zelenko", "Zuniga", "Ashworth", "Brannigan", "Cordero",
]
# Real NYC streets, the ZIP of the stretch the house numbers sit on, and the borough as City.
STREETS = [
    ("Grand Concourse", "10458", "Bronx", 2300, 2700), ("Fordham Rd", "10458", "Bronx", 300, 600),
    ("Southern Blvd", "10459", "Bronx", 900, 1200), ("Westchester Ave", "10462", "Bronx", 1900, 2200),
    ("Joralemon St", "11201", "Brooklyn", 100, 300), ("Atlantic Ave", "11217", "Brooklyn", 300, 600),
    ("Flatbush Ave", "11226", "Brooklyn", 2000, 2400), ("Nostrand Ave", "11216", "Brooklyn", 200, 400),
    ("Myrtle Ave", "11237", "Brooklyn", 1300, 1600), ("Ocean Pkwy", "11218", "Brooklyn", 300, 600),
    ("Broadway", "10027", "New York", 3000, 3200), ("Amsterdam Ave", "10024", "New York", 400, 600),
    ("Lexington Ave", "10029", "New York", 1600, 1800), ("St Nicholas Ave", "10032", "New York", 400, 600),
    ("Columbus Ave", "10025", "New York", 800, 1000), ("2nd Ave", "10003", "New York", 100, 300),
    ("Queens Blvd", "11375", "Queens", 10000, 11000), ("Jamaica Ave", "11432", "Queens", 16000, 17000),
    ("Roosevelt Ave", "11372", "Queens", 8000, 8600), ("Steinway St", "11103", "Queens", 3000, 3300),
    ("Victory Blvd", "10301", "Staten Island", 100, 600), ("Richmond Ave", "10314", "Staten Island", 1500, 2500),
    ("Forest Ave", "10310", "Staten Island", 400, 900),
]
UNITS = ["Apt 2B", "Apt 4C", "Apt 1F", "Apt 3", "Apt 5D", "#2R", "Fl 2", "Apt 6A", "Unit 3", "Apt 7G"]

PROGRAMS = {
    CSV: ["Workforce Development", "Housing Assistance", "Youth Services"],
    XLSX: ["Summer Youth Employment Program", "Housing Assistance", "Older Adult Services"],
    TSV: ["Workforce Dev.", "Housing Assistance", "Older Adult Services"],
    JSON_: ["SYEP", "Youth Services", "Older Adult Services"],
    XML: [],
}
# Unconfirmed proposals count their names separately (P-PROG-3). The planned decisions confirm these two groups.
PROGRAM_GROUPS = {"Workforce Dev.": "Workforce Development", "SYEP": "Summer Youth Employment Program"}
PROGRAM_PROPOSALS = [["Workforce Dev.", "Workforce Development", "same_name"], ["SYEP", "Summer Youth Employment Program", "initials"]]


@dataclass
class Address:
    house: str = ""
    street: str = ""
    unit: str = ""
    city: str = ""
    state: str = "NY"
    zip: str = ""


@dataclass
class Person:
    pid: int
    given: str
    surname: str
    middle: str
    dob: dt.date | None
    address: Address


@dataclass
class Rec:
    person: Person
    file: str
    client_id: str
    given: str
    surname: str
    middle: str
    dob: dt.date | None
    address: Address
    program: str
    reasons: list = field(default_factory=list)      # planned decision reason IDs at standardize (M04)
    notes: list = field(default_factory=list)        # planned audit notes at standardize (M04), a subset of what the row gets
    status: str = "counted"                          # counted, not_counted_yet or set_aside, before copies and identity
    row: int = 0                                     # file row number, filled in when written


class Plan:
    def __init__(self, seed):
        self.rng = random.Random(seed)
        self.pid = 0
        self.dobs = set()
        self.addresses = set()
        self.name_pairs = set()
        self.given = list(GIVEN)
        self.surnames = list(SURNAME)
        self.rng.shuffle(self.given)
        self.rng.shuffle(self.surnames)
        self.records = {f: [] for f in FILES}
        self.ids = {CSV: 1000, XLSX: 2000, TSV: 9000, JSON_: 3000, XML: 40000}
        self.copy_groups, self.conflicts, self.same_system, self.four_exact, self.likely, self.possible = [], [], [], [], [], []

    # people

    def dob(self):
        while True:
            d = dt.date(1940, 1, 1) + dt.timedelta(days=self.rng.randrange(0, 365 * 68))
            if d not in self.dobs:
                self.dobs.add(d)
                return d

    def address(self, unit_chance=0.4):
        while True:
            street, zip_, city, lo, hi = self.rng.choice(STREETS)
            house = str(self.rng.randrange(lo, hi) // 2 * 2 + self.rng.randrange(0, 2))
            if (house, street) in self.addresses:
                continue
            self.addresses.add((house, street))
            unit = self.rng.choice(UNITS) if self.rng.random() < unit_chance else ""
            return Address(house, street, unit, city, "NY", zip_)

    def person(self, given=None, surname=None, dob=True, address=None):
        while True:
            g = given or self.rng.choice(self.given)
            s = surname or self.rng.choice(self.surnames)
            if (g, s) not in self.name_pairs:
                break
            if given and surname:
                raise ValueError("planned name used twice")
        self.name_pairs.add((g, s))
        self.pid += 1
        middle = self.rng.choice("ABCDEFGHJKLMNPRSTW") if self.rng.random() < 0.3 else ""
        return Person(self.pid, g, s, middle, self.dob() if dob else None, address or self.address())

    def next_id(self, f):
        self.ids[f] += 1
        prefix = {CSV: "C-", XLSX: "H", TSV: "U", JSON_: "P-", XML: ""}[f]
        return f"{prefix}{self.ids[f]}"

    def rec(self, f, p, client_id=True, dob=True, program=None, **over):
        r = Rec(person=p, file=f, client_id=self.next_id(f) if client_id is True else (client_id or ""), given=p.given,
                surname=p.surname, middle=p.middle, dob=p.dob if dob else None, address=p.address,
                program=program if program is not None else (self.rng.choice(PROGRAMS[f]) if PROGRAMS[f] else ""))
        for k, v in over.items():
            setattr(r, k, v)
        self.records[f].append(r)
        return r


def build_plan(seed):
    pl = Plan(seed)
    rng = pl.rng

    # People in two or three files. These meet four_exact_facts: same DOB, surname, given name and address key (P-ID-3).
    for files in ((CSV, JSON_), (XLSX, XML), (CSV, TSV, XML)):
        p = pl.person(address=pl.address(unit_chance=0))
        pl.four_exact.append([pl.rec(f, p) for f in files])

    # Likely proposals: a nickname, and two birth-date digits swapped.
    p = pl.person(given="Robert", surname="Kessler")
    pl.likely.append([pl.rec(XLSX, p), pl.rec(TSV, p, given="Bob", address=pl.address())])
    p = pl.person(given="Maria", surname="Delgado")
    while True:
        first, swapped = p.dob.replace(day=12), p.dob.replace(day=21)
        if first not in pl.dobs and swapped not in pl.dobs:
            break
        p.dob = pl.dob()
    p.dob = first
    pl.dobs.update({first, swapped})
    pl.likely.append([pl.rec(JSON_, p), pl.rec(XML, p, dob=False)])
    pl.likely[-1][1].dob = swapped

    # Possible proposal: a rare name, no date of birth, who moved between the two files.
    p = pl.person(given="Thaddeus", surname="Brightwater", dob=False)
    pl.possible.append([pl.rec(CSV, p, dob=False), pl.rec(XML, p, dob=False, address=pl.address())])

    # Inside one file: an identical copy (no Client_ID, so name and DOB group it) and a possible copy under a nickname.
    p = pl.person(given="Lillian", surname="Ostrowski")
    a = pl.rec(XLSX, p, client_id=False, program="Housing Assistance")
    b = pl.rec(XLSX, p, client_id=False, program="Housing Assistance")
    pl.copy_groups.append(("identical", [a, b]))
    p = pl.person(given="William", surname="Ortiz")
    a = pl.rec(TSV, p, program="Older Adult Services")
    b = pl.rec(TSV, p, given="Bill", program="Older Adult Services")
    pl.copy_groups.append(("possible", [a, b]))

    # One Client_ID used for two different people in one system (P-DUP-3): not counted yet until decided.
    shared = pl.next_id(CSV)
    a = pl.rec(CSV, pl.person(), client_id=shared)
    b = pl.rec(CSV, pl.person(), client_id=shared)
    pl.conflicts.append([a, b])

    # One person enrolled twice under one ID in two programs: a same_system_id join, and two Goal 2 pairs.
    p = pl.person()
    uid = pl.next_id(TSV)
    pl.same_system.append([pl.rec(TSV, p, client_id=uid, program="Housing Assistance"),
                           pl.rec(TSV, p, client_id=uid, program="Workforce Dev.")])

    # A blank first or last name holds the count (name_missing).
    pl.rec(CSV, pl.person(), given="", reasons=["name_missing"], status="not_counted_yet")
    pl.rec(XLSX, pl.person(), given="", reasons=["name_missing"], status="not_counted_yet")
    pl.rec(JSON_, pl.person(), surname="", reasons=["name_missing"], status="not_counted_yet")

    # Twelve rows with neither a Client_ID nor a date of birth (no_identifier), spread over the files.
    for f, n in ((CSV, 3), (XLSX, 3), (TSV, 2), (JSON_, 2), (XML, 2)):
        for _ in range(n):
            pl.rec(f, pl.person(dob=False), client_id=False, dob=False, reasons=["no_identifier"], status="not_counted_yet")

    # Addresses: a placeholder, no fixed address, outside NYC, a ZIP+4, a ZIP with letters, and a blank ZIP a lookup can fill.
    # The placeholder street leaves the city, ZIP and state blank too: three missing parts and an incomplete address.
    pl.rec(JSON_, pl.person(address=Address("", "Unknown", "", "", "", "")),
           reasons=["address_incomplete", "address_part_missing", "placeholder_in_value"])
    pl.rec(XLSX, pl.person(address=Address("", "Homeless", "", "", "", "")), notes=["no_fixed_address"])
    pl.rec(TSV, pl.person(address=Address("120", "Newark Ave", "", "Jersey City", "NJ", "07302")),
           notes=["state_outside_nyc", "zip_outside_nyc"])
    pl.rec(XML, pl.person(address=Address("3080", "Broadway", "", "New York", "NY", "10027-6912")), notes=["zip_plus4_reduced"])
    pl.rec(JSON_, pl.person(address=Address("412", "Amsterdam Ave", "", "New York", "NY", "1002A")), reasons=["zip_not_valid"])
    pl.rec(XLSX, pl.person(address=Address("209", "Joralemon St", "", "Brooklyn", "NY", "")), reasons=["address_part_missing"])

    # Accented names, for the Windows Western callout on the TSV.
    for g, s in (("José", "Peña"), ("Zoë", "Müller"), ("Renée", "Gaspard"), ("Iñaki", "Muñoz"), ("Lucía", "Nowak")):
        pl.rec(TSV, pl.person(given=g, surname=s))

    # Fill each file to its planned size with ordinary people.
    for f, n in ((CSV, 60), (XLSX, 66), (TSV, 55), (JSON_, 55), (XML, 50)):
        while len(pl.records[f]) < n:
            pl.rec(f, pl.person())
        rng.shuffle(pl.records[f])
    return pl


# The practice reference and the location split (M05)

def zip5(a):
    return a.zip[:5]


def location_kind(r):
    """The plan's own reading of M04's location status: not_mapped with a reason, or to_place."""
    a = r.address
    if not a.house and a.street == "Homeless":
        return "no_fixed_address"
    if a.state and a.state != "NY":
        return "outside_nyc"
    if LOCATION_REASONS & set(r.reasons):
        return "address_problem"
    return "to_place"


def all_records(pl):
    return [r for f in FILES for r in pl.records[f]]


def plan_not_found(pl, seed):
    """A few addresses that only one row uses, left out of the reference so the example shows "not found"."""
    uses = {}
    for r in all_records(pl):
        if location_kind(r) == "to_place":
            uses.setdefault((r.address.house, r.address.street, zip5(r.address)), []).append(r)
    single = sorted(k for k, rs in uses.items() if len(rs) == 1 and rs[0].status == "counted")
    return set(random.Random(seed + 2).sample(single, NOT_FOUND))


def reference_rows(pl, seed, not_found):
    rng = random.Random(seed + 3)
    keys = sorted({(r.address.house, r.address.street, zip5(r.address), r.address.city) for r in all_records(pl)
                   if location_kind(r) == "to_place"})
    rows = []
    for house, street, z, city in keys:
        if (house, street, z) in not_found:
            continue
        lat0, lon0 = BOROUGH_BOX[city]
        rows.append([house, street, z, city, f"{lat0 + rng.uniform(-0.02, 0.02):.6f}", f"{lon0 + rng.uniform(-0.02, 0.02):.6f}"])
    return rows


def write_reference(rows, out):
    buf = io.StringIO()
    buf.write("#wholecount-sample\n")
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["house_number", "street", "zip", "borough", "latitude", "longitude"])
    w.writerows(rows)
    (out / REFERENCE).write_bytes(buf.getvalue().encode("utf-8"))


def _split(recs, outcome):
    split = {"mapped_high": 0, "mapped_low": 0, "not_mapped": {k: 0 for k in LOCATION_KINDS}}
    for r in recs:
        if r.status != "counted":
            continue
        kind = outcome(r)
        if kind == "mapped_high":
            split["mapped_high"] += 1
        else:
            split["not_mapped"][kind] += 1
    return split


def locations_section(pl, not_found, reference_count):
    def practice(r):
        k = location_kind(r)
        if k != "to_place":
            return k
        return "not_found" if (r.address.house, r.address.street, zip5(r.address)) in not_found else "mapped_high"

    def no_backend(r):
        k = location_kind(r)
        return "no_geocoder" if k == "to_place" else k

    looked_up = [r for r in all_records(pl) if location_kind(r) == "to_place"]
    return {
        "note": "Counted rows only, at standardize. Rows not counted yet are looked up too, so looked_up covers every row.",
        "reference": f"wholecount/sample_data/{REFERENCE}",
        "reference_rows": reference_count,
        "planned_not_found_rows": sorted(ref(r) for r in looked_up if practice(r) == "not_found"),
        "looked_up": len(looked_up),
        "practice_run": {f: _split(pl.records[f], practice) for f in FILES} | {"total": _split(all_records(pl), practice)},
        "no_backend_run": {f: _split(pl.records[f], no_backend) for f in FILES} | {"total": _split(all_records(pl), no_backend)},
        # M15: the split once the planned decisions are made, for the guide's second balance. The copies a person keeps one
        # of are set aside, and the same-ID conflict rows are counted once decided as different people.
        "practice_run_after_decisions": _split([r for r in all_records(pl)
                                                if not any(r is rs[1] for _, rs in pl.copy_groups)], practice),
    }


def reference_hash(out):
    return hashlib.sha256((out / REFERENCE).read_bytes()).hexdigest()


def hash_registered(digest):
    text = REFERENCE_PY.read_text(encoding="utf-8")
    block = re.search(r"SAMPLE_REFERENCE_HASHES = frozenset\(\{(.*?)\}\)", text, re.S)
    return bool(block) and digest in re.findall(r'"([0-9a-f]{64})"', block.group(1))


# Writing the files

def jammed_name(r):
    first = f"{r.given} {r.middle}".strip() if r.given else ""
    return f"{r.surname}, {first}".strip() if r.surname else first


def jammed_address(a):
    if not a.house and not a.city:
        return a.street
    parts = [a.house, a.street, a.unit, a.city, a.state, a.zip]
    return " ".join(p for p in parts if p)


def us_date(d):
    return f"{d.month:02d}/{d.day:02d}/{d.year}" if d else ""


def short_us_date(d):
    return f"{d.month}/{d.day}/{d.year}" if d else ""


def iso(d):
    return d.isoformat() if d else ""


def write_csv(pl, out):
    buf = io.StringIO(newline="")
    w = csv.writer(buf)
    w.writerow(["Client_ID", "Name", "Full_Address", "DOB", "Program"])
    for i, r in enumerate(pl.records[CSV], start=2):
        r.row = i
        w.writerow([r.client_id, jammed_name(r), jammed_address(r.address), us_date(r.dob), r.program])
    (out / CSV).write_bytes(buf.getvalue().encode("utf-8"))


XLSX_HEADER = ["Client ID", "First Name", "Last Name", "MI", "Address", "Apt", "City", "State", "ZIP", "Date of Birth", "Program"]
XLSX_TITLES = ["Housing program client roster", "Exported for the quarterly count by the housing team"]
NOTES_LINES = ["Rows are one enrollment each.", "Program names follow the funder's list.", "Questions go to the data team."]


def write_xlsx(pl, out):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clients"
    ws.append([XLSX_TITLES[0]])
    ws.append([XLSX_TITLES[1]])
    ws.append(XLSX_HEADER)
    for i, r in enumerate(pl.records[XLSX], start=4):
        r.row = i
        a = r.address
        ws.append([r.client_id or None, r.given or None, r.surname or None, r.middle or None,
                   f"{a.house} {a.street}".strip() or None, a.unit or None, a.city or None, a.state or None,
                   int(a.zip) if a.zip.isdigit() else (a.zip or None), r.dob, r.program or None])
        if r.dob:
            ws.cell(row=i, column=10).number_format = "mm/dd/yyyy"
    total_row = 4 + len(pl.records[XLSX])
    ws.append(["Total"] + [None] * 9 + [len(pl.records[XLSX])])
    notes = wb.create_sheet("Notes")
    for line in NOTES_LINES:
        notes.append([line])
    fixed = dt.datetime(2026, 10, 30)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = "make_example_data"
    wb.save(out / XLSX)
    return total_row


def write_tsv(pl, out, rng):
    buf = io.StringIO(newline="")
    w = csv.writer(buf, delimiter="\t")
    w.writerow(["uid", "first", "last", "SSN", "dob", "address_line", "apt", "city", "st", "zip", "program"])
    for i, r in enumerate(pl.records[TSV], start=2):
        r.row = i
        a = r.address
        ssn = f"9{rng.randrange(0, 100):02d}-{rng.randrange(10, 100)}-{rng.randrange(1000, 10000)}"
        w.writerow([r.client_id, r.given, r.surname, ssn, short_us_date(r.dob), f"{a.house} {a.street}".strip(), a.unit, a.city,
                    a.state, a.zip, r.program])
    (out / TSV).write_bytes(buf.getvalue().encode("cp1252"))


def write_json(pl, out):
    clients = []
    for i, r in enumerate(pl.records[JSON_], start=1):
        r.row = i
        a = r.address
        clients.append({"id": r.client_id, "first_name": r.given, "last_name": r.surname, "birth_date": iso(r.dob),
                        "address": {"street": f"{a.house} {a.street}".strip(), "unit": a.unit, "city": a.city, "state": a.state,
                                    "zip": a.zip},
                        "program": r.program})
    doc = {"export": {"generated": "2026-10-01", "source": "client portal", "clients": clients}}
    (out / JSON_).write_bytes((json.dumps(doc, indent=2, ensure_ascii=False) + "\n").encode("utf-8"))


def write_xml(pl, out):
    root = etree.Element("clients")
    for i, r in enumerate(pl.records[XML], start=1):
        r.row = i
        a = r.address
        c = etree.SubElement(root, "client")
        for tag, value in (("ClientNumber", r.client_id), ("FirstName", r.given), ("LastName", r.surname),
                           ("MiddleInitial", r.middle), ("BirthDate", iso(r.dob)), ("Street", f"{a.house} {a.street}".strip()),
                           ("Apt", a.unit), ("City", a.city), ("State", a.state), ("Zip", a.zip)):
            etree.SubElement(c, tag).text = value
    data = etree.tostring(root, pretty_print=True, xml_declaration=True, encoding="UTF-8")
    (out / XML).write_bytes(data)


# The expected numbers, from the plan

def ref(r):
    return [r.file, r.row]


def files_section(pl, total_row):
    n = {f: len(pl.records[f]) for f in FILES}
    plain = {"tables_found": 1, "tables_included": 1, "rows_excluded": 0, "rows_non_table": 0, "rows_above_header": 0,
             "rows_shared": 0, "balanced": True, "header_warning": None, "suggested_header_row": None}
    out = {
        CSV: dict(plain, format="csv", encoding="utf-8", rows_included=n[CSV], rows_header=1, non_empty_rows=n[CSV] + 1,
                  sensitive_dropped=[], tables=[{"sheet": None, "kind": "records", "included": True, "header_row": 1,
                                                 "rows": n[CSV]}]),
        TSV: dict(plain, format="tsv", encoding="windows-1252", rows_included=n[TSV], rows_header=1, non_empty_rows=n[TSV] + 1,
                  sensitive_dropped=["SSN"], tables=[{"sheet": None, "kind": "records", "included": True, "header_row": 1,
                                                      "rows": n[TSV]}]),
        JSON_: dict(plain, format="json", encoding=None, rows_included=n[JSON_], rows_header=0, non_empty_rows=n[JSON_],
                    sensitive_dropped=[], tables=[{"sheet": None, "kind": "records", "included": True, "header_row": 0,
                                                   "rows": n[JSON_]}]),
        XML: dict(plain, format="xml", encoding=None, rows_included=n[XML], rows_header=0, non_empty_rows=n[XML],
                  sensitive_dropped=[], tables=[{"sheet": None, "kind": "records", "included": True, "header_row": 0,
                                                 "rows": n[XML]}]),
    }
    clients_rows = len(XLSX_TITLES) + 1 + n[XLSX] + 1
    out[XLSX] = {
        "format": "xlsx", "encoding": None, "tables_found": 2, "tables_included": 0, "rows_included": 0,
        "rows_excluded": clients_rows + len(NOTES_LINES), "rows_non_table": 0, "rows_header": 0, "rows_above_header": 0,
        "rows_shared": 0, "non_empty_rows": clients_rows + len(NOTES_LINES), "balanced": True, "sensitive_dropped": [],
        "header_warning": "no_header_but_row_looks_like_names", "suggested_header_row": len(XLSX_TITLES) + 1,
        "tables": [
            {"sheet": "Clients", "kind": "unknown", "included": False, "header_row": 0, "rows": clients_rows},
            {"sheet": "Notes", "kind": "unknown", "included": False, "header_row": 0, "rows": len(NOTES_LINES)},
        ],
        "after_header_choice": {
            "sheet": "Clients", "header_row": len(XLSX_TITLES) + 1, "kind": "records", "included": True,
            "rows_included": n[XLSX] + 1, "rows_excluded": len(NOTES_LINES), "rows_header": 1, "rows_above_header": len(XLSX_TITLES),
            "total_row": total_row, "balanced": True,
        },
    }
    return out


def map_section():
    """The auto-mapping before any tidy-up step (M02), and after the proposed steps (M03). Blockers name fields only."""
    auto = {
        CSV: {"sheet": None, "fields": {"Client_ID": "Client_ID", "Address_1": "Full_Address", "Date_of_Birth": "DOB",
                                        "Program": "Program"},
              "blockers": ["First_Name", "Last_Name"]},
        XLSX: {"sheet": "Clients", "header_row": len(XLSX_TITLES) + 1,
               "fields": {"Client_ID": "Client ID", "First_Name": "First Name", "Last_Name": "Last Name", "Middle_Initial": "MI",
                          "Address_1": "Address", "Address_2": "Apt", "City": "City", "State": "State", "Zip": "ZIP",
                          "Date_of_Birth": "Date of Birth", "Program": "Program"}, "blockers": []},
        TSV: {"sheet": None, "fields": {"Client_ID": "uid", "First_Name": "first", "Last_Name": "last", "Date_of_Birth": "dob",
                                        "Address_1": "address_line", "Address_2": "apt", "City": "city", "State": "st", "Zip": "zip",
                                        "Program": "program"}, "blockers": []},
        JSON_: {"sheet": None, "fields": {"Client_ID": "id", "First_Name": "first_name", "Last_Name": "last_name",
                                          "Date_of_Birth": "birth_date", "Address_1": "address.street", "Address_2": "address.unit",
                                          "City": "address.city", "State": "address.state", "Zip": "address.zip", "Program": "program"},
                "blockers": []},
        XML: {"sheet": None, "fields": {"Client_ID": "ClientNumber", "First_Name": "FirstName", "Last_Name": "LastName",
                                        "Middle_Initial": "MiddleInitial", "Date_of_Birth": "BirthDate", "Address_1": "Street",
                                        "Address_2": "Apt", "City": "City", "State": "State", "Zip": "Zip"}, "blockers": []},
    }
    after = {f: dict(v) for f, v in auto.items()}
    after[CSV] = {
        "sheet": None,
        "steps": [{"kind": "split_address", "column": "Full_Address"}, {"kind": "split_name", "column": "Name"}],
        "fields": {"Client_ID": "Client_ID", "Date_of_Birth": "DOB", "Program": "Program", "First_Name": "Name_first",
                   "Middle_Initial": "Name_middle", "Last_Name": "Name_last", "Address_1": "Full_Address_street",
                   "Address_2": "Full_Address_address2", "City": "Full_Address_city", "State": "Full_Address_state",
                   "Zip": "Full_Address_zip"},
        "blockers": [],
    }
    for f in (XLSX, TSV, JSON_, XML):
        after[f]["steps"] = []
    return {"auto": auto, "after_steps": after}


def standardize_section(pl, total_row):
    out = {}
    for f in FILES:
        recs = pl.records[f]
        counted = sum(1 for r in recs if r.status == "counted")
        nc = sum(1 for r in recs if r.status == "not_counted_yet")
        problems = [{"row": r.row, "reasons": sorted(set(r.reasons))} for r in recs if r.reasons]
        notes = [{"row": r.row, "notes": sorted(r.notes)} for r in recs if r.notes]
        set_aside = []
        if f == XLSX:
            set_aside = [{"row": total_row, "reason": "summary_row"}]
        out[f] = {"rows_received": len(recs) + len(set_aside), "counted": counted, "not_counted_yet": nc,
                  "set_aside": len(set_aside), "planned_problem_rows": sorted(problems, key=lambda p: p["row"]),
                  "planned_set_aside_rows": set_aside, "planned_notes": sorted(notes, key=lambda p: p["row"])}
    out["no_identifier_total"] = sum(1 for f in FILES for r in pl.records[f] if "no_identifier" in r.reasons)
    return out


def people_section(pl):
    return {
        "copy_groups": [{"kind": kind, "rows": [ref(r) for r in rs]} for kind, rs in pl.copy_groups],
        "same_id_conflicts": [{"client_id": rs[0].client_id, "rows": [ref(r) for r in rs]} for rs in pl.conflicts],
        "auto_joins": {
            "same_system_id": [[ref(r) for r in rs] for rs in pl.same_system],
            "four_exact_facts": [[ref(r) for r in rs] for rs in pl.four_exact],
            "earlier_batch": [],
        },
        "proposals": {
            "likely": [{"rows": [ref(r) for r in rs], "why": why} for rs, why in zip(pl.likely, ["nickname", "dob_digits_swapped"])],
            "possible": [{"rows": [ref(r) for r in rs], "why": "rare_name_moved"} for rs in pl.possible],
        },
    }


def programs_section(pl):
    names = sorted({r.program for f in FILES for r in pl.records[f] if r.program})
    no_program = [f for f in FILES if not PROGRAMS[f]]
    return {"names": names, "proposals": [{"names": [a, b], "reason": why} for a, b, why in PROGRAM_PROPOSALS],
            "not_in_goal_2": [{"file": f, "rows": len(pl.records[f])} for f in no_program]}


def decisions_section(pl):
    return {
        "header_rows": [{"file": XLSX, "sheet": "Clients", "header_row": len(XLSX_TITLES) + 1}],
        "program_groups": [{"names": [a, b], "decision": "confirm", "label": b} for a, b, _ in PROGRAM_PROPOSALS],
        "copy_groups": [{"rows": [ref(r) for r in rs], "decision": "keep_one", "keep": ref(rs[0])} for _, rs in pl.copy_groups],
        "same_id_conflicts": [{"rows": [ref(r) for r in rs], "decision": "different_people"} for rs in pl.conflicts],
        "pairs": [{"rows": [ref(r) for r in rs], "decision": "same_person"} for rs in pl.likely + pl.possible],
    }


class Groups:
    def __init__(self):
        self.parent = {}

    def find(self, x):
        self.parent.setdefault(x, x)
        while self.parent[x] != x:
            self.parent[x] = self.parent[self.parent[x]]
            x = self.parent[x]
        return x

    def join(self, a, b):
        self.parent[self.find(a)] = self.find(b)


def goals(records, joins, programs_grouped):
    """(Goal 1, Goal 2, per program) for records joined into people by the given pairs of records."""
    g = Groups()
    for r in records:
        g.find(id(r))
    for a, b in joins:
        if id(a) in g.parent and id(b) in g.parent:
            g.join(id(a), id(b))
    masters = {g.find(id(r)) for r in records}
    pairs, per_program = set(), {}
    for r in records:
        if not r.program:
            continue
        name = PROGRAM_GROUPS.get(r.program, r.program) if programs_grouped else r.program
        pairs.add((g.find(id(r)), name))
    for _, name in pairs:
        per_program[name] = per_program.get(name, 0) + 1
    return len(masters), len(pairs), dict(sorted(per_program.items()))


def chain(rs):
    return [(rs[i], rs[i + 1]) for i in range(len(rs) - 1)]


def counts_section(pl, standardize):
    all_recs = [r for f in FILES for r in pl.records[f]]
    received = sum(standardize[f]["rows_received"] for f in FILES)
    conflict_rows = [r for rs in pl.conflicts for r in rs]
    counted_before = [r for r in all_recs if r.status == "counted" and r not in conflict_rows]
    auto = [p for rs in pl.same_system + pl.four_exact for p in chain(rs)]
    open_merges = [tuple(rs) for _, rs in pl.copy_groups] + [tuple(rs) for rs in pl.likely + pl.possible]
    g1_high, g2_high, _ = goals(counted_before, auto, False)
    g1_low, g2_low, _ = goals(counted_before, auto + open_merges, False)
    set_aside_before = 1
    before = {"rows_received": received, "counted": len(counted_before),
              "not_counted_yet": received - len(counted_before) - set_aside_before, "set_aside": set_aside_before,
              "goal_1": {"low": g1_low, "high": g1_high}, "goal_2": {"low": g2_low, "high": g2_high},
              "open_pairs": len(pl.likely) + len(pl.possible), "open_copy_groups": len(pl.copy_groups)}

    kept_aside = [rs[1] for _, rs in pl.copy_groups]
    counted_after = [r for r in all_recs if r.status == "counted" and r not in kept_aside]
    decided = [tuple(rs) for rs in pl.likely + pl.possible]
    g1, g2, per_program = goals(counted_after, auto + decided, True)
    set_aside_after = 1 + len(kept_aside)
    after = {"rows_received": received, "counted": len(counted_after),
             "not_counted_yet": received - len(counted_after) - set_aside_after, "set_aside": set_aside_after,
             "goal_1": {"low": g1, "high": g1}, "goal_2": {"low": g2, "high": g2}, "per_program": per_program,
             "open_pairs": 0, "open_copy_groups": 0}
    return {"before_decisions": before, "after_decisions": after}


def outputs_section(pl, counts):
    """Row counts per file after the planned decisions. The review file holds every row with an open problem: the rows not
    counted yet, and the counted rows whose planned problems don't hold the count (location and worth-a-look reasons)."""
    a = counts["after_decisions"]
    counted_yes = sum(1 for f in FILES for r in pl.records[f] if r.status == "counted" and r.reasons)
    run = {"Final_Clean_Mapped": a["counted"], "Set_Aside_Records": a["set_aside"],
           "Manual_Review_Needed_counted_no": a["not_counted_yet"], "Manual_Review_Needed_counted_yes": counted_yes,
           "Manual_Review_Needed": a["not_counted_yet"] + counted_yes, "Unique_Clients": a["goal_1"]["high"],
           "Unique_Clients_By_Program": a["goal_2"]["high"], "Data_Quality_Report": len(FILES)}
    return {"practice_run": dict(run), "no_geocoder_run": dict(run),
            "note": "One table per file. Map files aren't listed: a practice run writes none, and the example data has no districts."}


def expected(pl, total_row, locations):
    std = standardize_section(pl, total_row)
    counts = counts_section(pl, std)
    return {
        "seed": SEED, "plan": "tools/make_example_data.py",
        "files": files_section(pl, total_row),
        "map": map_section(),
        "standardize": std,
        "people": people_section(pl),
        "programs": programs_section(pl),
        "decisions": decisions_section(pl),
        "counts": counts,
        "locations": locations,
        "outputs": outputs_section(pl, counts),
    }


def generate(seed, out):
    out.mkdir(parents=True, exist_ok=True)
    pl = build_plan(seed)
    ssn_rng = random.Random(seed + 1)
    write_csv(pl, out)
    total_row = _write_xlsx_if_changed(pl, out)
    write_tsv(pl, out, ssn_rng)
    write_json(pl, out)
    write_xml(pl, out)
    not_found = plan_not_found(pl, seed)
    rows = reference_rows(pl, seed, not_found)
    write_reference(rows, out)
    exp = expected(pl, total_row, locations_section(pl, not_found, len(rows)))
    exp["seed"] = seed
    (out / "EXPECTED.json").write_bytes((json.dumps(exp, indent=2, ensure_ascii=False) + "\n").encode("utf-8"))
    return exp


def _write_xlsx_if_changed(pl, out):
    """openpyxl stamps times into the zip, so an unchanged workbook is kept as it is rather than rewritten."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp = pathlib.Path(tmp)
        total_row = write_xlsx(pl, tmp)
        target = out / XLSX
        if not target.exists() or _cells(target) != _cells(tmp / XLSX):
            target.write_bytes((tmp / XLSX).read_bytes())
    return total_row


def _cells(path):
    wb = openpyxl.load_workbook(path)
    try:
        return [(ws.title, [[c.value for c in row] for row in ws.iter_rows()]) for ws in wb.worksheets]
    finally:
        wb.close()


def check(seed, out):
    """Regenerate into a temporary folder and compare. Prints counts only."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp = pathlib.Path(tmp)
        generate(seed, tmp)
        differing = 0
        for name in FILES + (REFERENCE, "EXPECTED.json"):
            a, b = out / name, tmp / name
            if not a.exists():
                differing += 1
            elif name.endswith(".xlsx"):
                differing += _cells(a) != _cells(b)
            else:
                differing += a.read_bytes() != b.read_bytes()
    registered = hash_registered(reference_hash(out)) if (out / REFERENCE).exists() else False
    print(f"files_compared {len(FILES) + 2}")
    print(f"files_differing {differing}")
    print(f"reference_hash_registered {int(registered)}")
    return 1 if differing or not registered else 0


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=SEED)
    ap.add_argument("--out", type=pathlib.Path, default=OUT)
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args(argv)
    if args.check:
        return check(args.seed, args.out)
    generate(args.seed, args.out)
    print(f"files_written {len(FILES) + 2}")
    digest = reference_hash(args.out)
    print(f"reference_sha256 {digest}")
    print(f"reference_hash_registered {int(hash_registered(digest))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
